from __future__ import annotations

import sys
from pathlib import Path
from decimal import Decimal
import csv
from io import BytesIO, StringIO


def read_xlsx(path: Path):
    from openpyxl import load_workbook  # type: ignore

    wb = load_workbook(filename=str(path), data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    try:
        idx_name = headers.index("name")
    except ValueError:
        idx_name = -1
    try:
        idx_desc = headers.index("description")
    except ValueError:
        idx_desc = -1
    price_candidates = [
        i for i, h in enumerate(headers) if h in {"price", "price_uah", "ціна", "tsina"}
    ]
    idx_price = price_candidates[0] if price_candidates else -1
    imported = 0
    skipped = 0
    for r in ws.iter_rows(min_row=2):
        name_cell = r[idx_name] if idx_name >= 0 else None
        price_cell = r[idx_price] if idx_price >= 0 else None
        desc_cell = r[idx_desc] if idx_desc >= 0 else None
        name = (
            str(name_cell.value).strip()
            if name_cell and name_cell.value is not None
            else ""
        )
        if not name:
            skipped += 1
            continue
        raw_price = price_cell.value if price_cell else None
        if raw_price is None or (
            isinstance(raw_price, str) and not str(raw_price).strip()
        ):
            skipped += 1
            continue
        try:
            if isinstance(raw_price, (int, float)):
                Decimal(str(raw_price))
            else:
                Decimal(str(raw_price).replace(",", ".").strip())
        except Exception:
            skipped += 1
            continue
        imported += 1
    return headers, imported, skipped, (ws.max_row - 1)


def read_csv(path: Path):
    text = path.read_text(encoding="utf-8-sig")
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=[",", ";", "\t", "|"])
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ";" if ";" in sample and "," not in sample else ","
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    rows = list(reader)
    headers = [str(h or "").strip().lower() for h in rows[0]] if rows else []
    try:
        idx_name = headers.index("name")
    except ValueError:
        idx_name = -1
    try:
        idx_desc = headers.index("description")
    except ValueError:
        idx_desc = -1
    price_candidates = [
        i for i, h in enumerate(headers) if h in {"price", "price_uah", "ціна", "tsina"}
    ]
    idx_price = price_candidates[0] if price_candidates else -1
    imported = 0
    skipped = 0
    for row in rows[1:]:
        name = (
            str(row[idx_name]).strip() if idx_name >= 0 and idx_name < len(row) else ""
        )
        if not name:
            skipped += 1
            continue
        raw_price_value = (
            row[idx_price] if idx_price >= 0 and idx_price < len(row) else None
        )
        if raw_price_value is None or (
            isinstance(raw_price_value, str) and not str(raw_price_value).strip()
        ):
            skipped += 1
            continue
        try:
            Decimal(str(raw_price_value).replace(",", ".").strip())
        except Exception:
            skipped += 1
            continue
        imported += 1
    return headers, imported, skipped, max(len(rows) - 1, 0), delimiter


def main():
    base = Path("agent_context_documents")
    csv_path = base / "boxcatering-menu.csv"
    xlsx_path = base / "boxcatering-menu.xlsx"
    if csv_path.exists():
        headers, imported, skipped, total, delim = read_csv(csv_path)
        print(f"CSV ({delim}) headers: {headers}")
        print(f"CSV imported={imported} skipped={skipped} total={total}")
    else:
        print("CSV file not found")
    if xlsx_path.exists():
        headers, imported, skipped, total = read_xlsx(xlsx_path)
        print(f"XLSX headers: {headers}")
        print(f"XLSX imported={imported} skipped={skipped} total={total}")
    else:
        print("XLSX file not found")


if __name__ == "__main__":
    main()
